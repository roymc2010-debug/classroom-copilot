import os
import io
import re
import json
import datetime
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from pypdf import PdfReader

# Permite relajar el scope en caso de discrepancias menores de Google OAuthlib
os.environ['OAUTHLIB_RELAX_TOKEN_SCOPE'] = '1'

SCOPES = [
    'https://www.googleapis.com/auth/classroom.courses.readonly',
    'https://www.googleapis.com/auth/classroom.coursework.me.readonly',
    'https://www.googleapis.com/auth/classroom.student-submissions.me.readonly',
    'https://www.googleapis.com/auth/classroom.courseworkmaterials.readonly',
    'https://www.googleapis.com/auth/classroom.announcements.readonly',
    'https://www.googleapis.com/auth/drive.readonly',
    'https://www.googleapis.com/auth/drive.file',
    'https://www.googleapis.com/auth/gmail.readonly',
    'https://www.googleapis.com/auth/userinfo.email'
]

PDF_CACHE_DIR = os.path.join("data", "pdf_cache")

def get_cached_attachment_data(file_id: str) -> dict:
    """Lee del almacenamiento en disco las consignas extraídas y el texto del documento para carga instantánea."""
    if not file_id:
        return {}
    cache_path = os.path.join(PDF_CACHE_DIR, f"{file_id}.json")
    if os.path.exists(cache_path):
        try:
            with open(cache_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_cached_attachment_data(file_id: str, data: dict):
    """Guarda en caché persistente en disco el análisis de consignas del archivo."""
    if not file_id or not data:
        return
    try:
        os.makedirs(PDF_CACHE_DIR, exist_ok=True)
        cache_path = os.path.join(PDF_CACHE_DIR, f"{file_id}.json")
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
    except Exception as e:
        print(f"Error guardando caché de PDF {file_id}: {e}")

def sanitize_extracted_text(text: str) -> str:
    """
    Sanitiza y reconstruye oraciones completas de textos extraídos de PDFs, diapositivas y documentos:
    - Normaliza retornos de carro (\r\n a \n).
    - Une palabras partidas por guiones de separación silábica al final de línea (ej. experi-\nmentos -> experimentos).
    - Une saltos de línea duros (\n): si un renglón no termina con un punto o puntuación de cierre
      y el siguiente no es una nueva viñeta o inciso, lo une con el siguiente renglón mediante un espacio.
    - Limpia dobles/múltiples espacios y espaciados anómalos antes de signos de puntuación.
    """
    if not text:
        return ""

    # 1. Normalizar saltos de línea
    t = text.replace('\r\n', '\n').replace('\r', '\n')

    # 2. Deshacer guiones de separación de palabras al final de línea
    t = re.sub(r'(\w+)[-–—]\s*\n\s*(\w+)', r'\1\2', t)

    # 3. Reconstruir oraciones completas causadas por saltos de línea duros
    raw_lines = t.split('\n')
    cleaned_lines = []
    current_line = ''
    CLOSING_PUNCT = ('.', '!', '?', ':', ';', '.)', '!"', '?"', '."', "'")
    bullet_re = re.compile(r'^(?:[-*•–—]|\d+[\.\)]|[a-zA-Z][\.\)])\s+')

    for line in raw_lines:
        s = line.strip()
        if not s:
            if current_line:
                cleaned_lines.append(current_line)
                current_line = ''
            continue

        if not current_line:
            current_line = s
        else:
            ends_with_punct = any(current_line.endswith(p) for p in CLOSING_PUNCT)
            is_new_bullet = bool(bullet_re.match(s))

            if ends_with_punct or is_new_bullet:
                cleaned_lines.append(current_line)
                current_line = s
            else:
                # Si el renglón no termina con punto o puntuación de cierre,
                # unirlo con el siguiente renglón con un espacio para reconstruir la oración completa.
                current_line = f"{current_line} {s}"

    if current_line:
        cleaned_lines.append(current_line)

    res = '\n'.join(cleaned_lines)

    # 4. Limpiar espacios dobles o múltiples y espaciado previo a signos
    res = re.sub(r'[ \t]{2,}', ' ', res)
    res = re.sub(r'\s+([.,;:!?])', r'\1', res)
    return res.strip()

def extract_actionable_consignas(text: str, course_name: str = "", task_title: str = "", max_items: int = 5) -> list[str]:
    """
    Extrae de forma robusta y rigurosa las consignas reales y requisitos operativos de una tarea universitaria.
    Aplica sanitización previa de saltos de línea duros, consulta al motor de IA (Groq/OpenRouter/Gemini/OpenAI)
    con el prompt de requisitos operativos críticos, y cuenta con un fallback heurístico exhaustivo
    anti-oraciones cortadas.
    """
    if not text:
        return []

    clean_text = sanitize_extracted_text(text)
    if not clean_text:
        return []

    # 1. Intentar extracción con modelo de IA y prompt especializado
    try:
        from services.ai_service import extract_task_consignas_ai_sync
        ai_actions = extract_task_consignas_ai_sync(clean_text, course_name=course_name, task_title=task_title)
        if ai_actions and len(ai_actions) >= 1:
            return ai_actions[:max_items]
    except Exception as e:
        print(f"[Classroom Service] Fallback heurístico en consignas: {e}")

    # 2. Fallback Heurístico Robusto (sin dependencias externas / offline)
    lines = [l.strip() for l in clean_text.split('\n') if l.strip()]

    IGNORE_PATTERNS = [
        r'universidad', r'departamento', r'divisi[oó]n', r'centro universitario', r'cucei',
        r'profesor\b', r'docente\b', r'alumno\b', r'estudiante\b', r'c[oó]digo:', r'fecha:',
        r'semestre\b', r'ciclo escolar', r'licenciatura', r'cr[eé]ditos',
        r'criterios? de evaluaci[oó]n', r'ponderaci[oó]n', r'r[uú]brica', r'bibliograf[ií]a',
        r'p[aá]gina \d+', r'^\d+\s*$'
    ]
    ignore_re = re.compile('|'.join(IGNORE_PATTERNS), re.IGNORECASE)

    ACTION_VERBS = [
        # Operativos / Organización / Modalidad / Lugar / Trámite
        r'formar(?:\s+equipos)?', r'integrar', r'agendar(?:\s+una\s+visita)?', r'visitar', r'visita',
        r'inspecci[oó]n', r'inspeccionar', r'investigaci[oó]n', r'investigar', r'revisi[oó]n',
        r'laboratorio', r'aula\s+[a-zA-Z0-9]+', r'citas?', r'integrantes',
        # Académicos / Técnicos
        r'resuelve', r'resolver', r'calcula', r'calcular', r'determina', r'determinar',
        r'obten(?:er|ga)?', r'halla(?:r)?', r'grafica(?:r)?', r'demuestra', r'demostrar',
        r'elabora(?:r)?', r'realiza(?:r)?', r'redacta(?:r)?', r'desarrolla(?:r)?',
        r'simula(?:r)?', r'analiza(?:r)?', r'compara(?:r)?',
        r'entrega(?:r)?', r'sube', r'subir', r'contesta(?:r)?', r'responde(?:r)?',
        r'encuentra', r'encontrar', r'aplica(?:r)?', r'ejercicios?', r'problemas?',
        r'transformada de laplace', r'funci[oó]n de transferencia', r'diagrama',
        r'ecuaci[oó]n', r'circuito', r'cuestionario', r'reporte', r'ensayo',
        r'plantilla\s+ieee', r'formato\s+pdf'
    ]
    action_re = re.compile(r'\b(' + '|'.join(ACTION_VERBS) + r')\b', re.IGNORECASE)

    prefix_clean_re = re.compile(r'^(?:instrucciones?|consigna|objetivo|actividad|tarea|ejercicio\s*\d*|problema\s*\d*)\s*[:.-]\s*', re.IGNORECASE)
    bullet_clean_re = re.compile(r'^(?:[-*•–—]|\d+[\.\)]|[a-zA-Z][\.\)])\s*')
    dangling_connector_re = re.compile(r'[\s,]+(?:y|e|o|u|de|en|con|para|que|a|al|del)$', re.IGNORECASE)
    incomplete_list_re = re.compile(r'(?:que\s+contenga\s+la\s+siguiente(?:\s+informaci[oó]n)?|los\s+siguientes\s+puntos|lo\s+siguiente)\s*[:.]?$', re.IGNORECASE)

    candidate_sentences = []

    for line in lines:
        if ignore_re.search(line) and not action_re.search(line):
            continue

        sentences = re.split(r'(?<=[.!?])\s+', line)
        for s in sentences:
            s_clean = s.strip()
            s_clean = bullet_clean_re.sub('', s_clean).strip()
            s_clean = prefix_clean_re.sub('', s_clean).strip()

            if len(s_clean) < 10:
                continue

            if course_name and s_clean.lower() == course_name.lower():
                continue
            if task_title and s_clean.lower() == task_title.lower():
                continue

            # Omitir si la frase es solo el nombre de la materia repetido
            if re.match(r'^(?:teor[ií]a de sistemas|c[aá]lculo|control|sistemas inteligentes)(?:\s+[ivx0-9a-z]+)*$', s_clean, re.IGNORECASE):
                continue

            if action_re.search(s_clean):
                # Arreglar listas incompletas o que terminan en 'contenga la siguiente información'
                if incomplete_list_re.search(s_clean):
                    s_clean = incomplete_list_re.sub(r' (consultar especificaciones en el documento)', s_clean).strip()

                # Limpiar conectores huérfanos al final de la oración
                while dangling_connector_re.search(s_clean):
                    s_clean = dangling_connector_re.sub('', s_clean).strip()

                if len(s_clean) < 12:
                    continue

                s_formatted = s_clean[0].upper() + s_clean[1:]
                if not s_formatted.endswith(('.', '!', '?')):
                    s_formatted += '.'
                if s_formatted not in candidate_sentences:
                    candidate_sentences.append(s_formatted)

    filtered = []
    for c in candidate_sentences:
        if len(c) > 300:
            sub = re.split(r'(?<=[.!?])\s+', c)
            c = sub[0]
            if not c.endswith('.'):
                c += '.'
        # Validar que no termine en conector trunco
        if re.search(r'\b(?:y|e|o|u|de|en|con|para|que|del)\.$', c, re.IGNORECASE):
            continue
        if len(c) >= 15 and c not in filtered:
            filtered.append(c)

    return filtered[:max_items]

def inject_authuser(url: str, user_email: str = None) -> str:
    """
    Inyecta deterministamente ?authuser={user_email} en enlaces externos de Google
    (Classroom, Drive, Docs, etc.) para que el navegador del alumno abra directamente
    su perfil escolar y evite errores de 'cuenta incorrecta / clase no encontrada'.
    """
    if not url or not user_email:
        return url or ""
    if "authuser=" in url:
        return url
    if "google.com" in url or "classroom." in url or "drive." in url or "docs." in url:
        # Si contiene /u/0/ o /u/1/ en URLs de Google, limpiarlo para evitar forzar la cuenta incorrecta
        clean_url = re.sub(r'/u/\d+/', '/', url)
        if "authuser=" in clean_url:
            return clean_url
        sep = "&" if "?" in clean_url else "?"
        return f"{clean_url}{sep}authuser={user_email}"
    return url

def get_credentials():
    creds = None
    if os.path.exists('token.json'):
        try:
            creds = Credentials.from_authorized_user_file('token.json', SCOPES)
        except Exception:
            creds = None
    return creds

def get_classroom_service(creds=None):
    c = creds if creds else get_credentials()
    return build('classroom', 'v1', credentials=c)

def get_drive_service(creds=None):
    c = creds if creds else get_credentials()
    return build('drive', 'v3', credentials=c)

def get_gmail_service(creds=None):
    try:
        c = creds if creds else get_credentials()
        return build('gmail', 'v1', credentials=c)
    except Exception:
        return None

def extract_pdf_text_from_drive(drive_service, file_id: str) -> str:
    """Extrae texto de un PDF o Documento en Drive bajo demanda (con sanitización de renglones y caché)."""
    if not file_id:
        return ""
    cached = get_cached_attachment_data(file_id)
    if cached and cached.get("text"):
        return sanitize_extracted_text(cached["text"])
    if not drive_service:
        return ""

    try:
        request = drive_service.files().get_media(fileId=file_id)
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()

        # 1. Intentar como PDF
        try:
            fh.seek(0)
            reader = PdfReader(fh)
            text_content = []
            for page in reader.pages[:10]:
                t = page.extract_text()
                if t:
                    text_content.append(t)
            raw = "\n".join(text_content).strip()
            res = sanitize_extracted_text(raw)
            if res:
                return res
        except Exception:
            pass

        # 2. Intentar como DOCX (Word)
        try:
            fh.seek(0)
            import docx
            doc = docx.Document(fh)
            text_content = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        text_content.append(row_text)
            raw = "\n".join(text_content).strip()
            res = sanitize_extracted_text(raw)
            if res:
                return res
        except Exception:
            pass

        # 3. Intentar como texto plano / markdown
        try:
            fh.seek(0)
            raw = fh.read().decode('utf-8', errors='ignore').strip()
            res = sanitize_extracted_text(raw)
            if res and len(res) > 20 and not res.startswith('\x00'):
                return res
        except Exception:
            pass

    except Exception as e:
        print(f"Aviso: descarga get_media de {file_id} ({e}), probando exportación alternativa...")

    # Intento de respaldo: si es un Google Doc nativo
    try:
        req_export = drive_service.files().export_media(fileId=file_id, mimeType='text/plain')
        fh_exp = io.BytesIO()
        downloader_exp = MediaIoBaseDownload(fh_exp, req_export)
        done = False
        while not done:
            _, done = downloader_exp.next_chunk()
        fh_exp.seek(0)
        raw = fh_exp.read().decode('utf-8', errors='ignore').strip()
        return sanitize_extracted_text(raw)
    except Exception as e_exp:
        print(f"Aviso: no se pudo extraer texto del archivo {file_id}: {e_exp}")
        return ""

def get_task_attachment_summary(
    drive_service,
    file_id: str,
    course_name: str = "",
    task_title: str = "",
    force_refresh: bool = False
) -> dict:
    """
    Obtiene el resumen y acciones del adjunto, aprovechando la caché en disco
    e invalidando automáticamente cualquier entrada con oraciones cortadas o frases truncadas.
    """
    if not file_id:
        return {"actions": [], "text": ""}
    cached = get_cached_attachment_data(file_id)
    if not force_refresh and cached and cached.get("actions"):
        # Auto-invalidar si la caché previa contiene oraciones cortadas por el bug antiguo
        has_broken_action = any(
            re.search(r'\b(?:y|e|o|u|de|en|con|para|que|del)\.$', str(a).strip(), re.IGNORECASE) or
            "la siguiente." in str(a).lower() or
            "los siguientes puntos." in str(a).lower() or
            len(str(a).strip()) < 10
            for a in cached["actions"]
        )
        if not has_broken_action:
            return cached

    text = extract_pdf_text_from_drive(drive_service, file_id)
    actions = extract_actionable_consignas(text, course_name=course_name, task_title=task_title)

    data = {
        "file_id": file_id,
        "actions": actions,
        "text": text[:3500]
    }
    if text or actions:
        save_cached_attachment_data(file_id, data)
    return data

def get_active_courses(service):
    """
    Obtiene todas las materias del usuario de forma ultra-resiliente.
    Evita filtros excesivamente restrictivos que causan listas vacías en cuentas institucionales o escolares.
    """
    courses = []
    page_token = None

    # Intento 1: Consulta amplia sin restricción de studentId ni courseStates
    try:
        while True:
            res = service.courses().list(pageSize=100, pageToken=page_token).execute()
            batch = res.get('courses', [])
            courses.extend(batch)
            page_token = res.get('nextPageToken')
            if not page_token:
                break
    except Exception as e:
        print(f"[Classroom API] Intento 1 (sin filtros) falló: {e}. Probando intento 2...")
        courses = []

    # Intento 2: Si el intento 1 devolvió 0 o falló, probar con studentId='me'
    if not courses:
        try:
            page_token = None
            while True:
                res = service.courses().list(studentId='me', pageSize=100, pageToken=page_token).execute()
                batch = res.get('courses', [])
                courses.extend(batch)
                page_token = res.get('nextPageToken')
                if not page_token:
                    break
        except Exception as e2:
            print(f"[Classroom API] Intento 2 (studentId='me') falló: {e2}")

    # Intento 3: Probar con studentId='me' y courseStates=['ACTIVE'] como último recurso
    if not courses:
        try:
            page_token = None
            while True:
                res = service.courses().list(studentId='me', courseStates=['ACTIVE'], pageSize=100, pageToken=page_token).execute()
                batch = res.get('courses', [])
                courses.extend(batch)
                page_token = res.get('nextPageToken')
                if not page_token:
                    break
        except Exception as e3:
            print(f"[Classroom API] Intento 3 falló: {e3}")

    # Filtrar únicamente si hay materias y descartar solo las archivadas si existen materias activas
    non_archived = [c for c in courses if c.get('courseState') not in ('ARCHIVED', 'DECLINED')]
    final_list = non_archived if non_archived else courses

    print(f"[Classroom API] Cursos recuperados ({len(final_list)}): {[c.get('name') for c in final_list]}")
    return final_list

def get_all_tasks(creds=None, user_email=None):
    service = get_classroom_service(creds=creds)

    # Si no se pasó user_email, intentar extraerlo del perfil
    if not user_email and creds:
        try:
            profile = service.userProfiles().get(userId='me').execute()
            user_email = profile.get('emailAddress', '').lower().strip()
        except Exception:
            pass

    courses = get_active_courses(service)
    tasks = []

    for course in courses:
        course_id = course['id']
        course_name = course.get('name', 'Materia').replace('_', ' ')

        try:
            cw_res = service.courses().courseWork().list(courseId=course_id).execute()
            course_works = cw_res.get('courseWork', [])
        except Exception as e:
            print(f"[Classroom API] Error listando tareas para materia {course_name} ({course_id}): {e}")
            course_works = []

        for cw in course_works:
            cw_id = cw['id']
            title = cw.get('title', 'Sin título')
            desc = cw.get('description', '')
            alt_link = inject_authuser(cw.get('alternateLink', ''), user_email)
            max_points = cw.get('maxPoints')

            due_date = cw.get('dueDate')
            due_time = cw.get('dueTime')
            due_date_iso = None

            if due_date:
                year = due_date.get('year')
                month = due_date.get('month')
                day = due_date.get('day')
                if due_time:
                    hour = due_time.get('hours', 0)
                    minute = due_time.get('minutes', 0)
                else:
                    hour = 23
                    minute = 59
                dt = datetime.datetime(year, month, day, hour, minute, tzinfo=datetime.timezone.utc)
                due_date_iso = dt.isoformat()

            classroom_status = 'PENDIENTE'
            assigned_grade = None

            try:
                sub_res = service.courses().courseWork().studentSubmissions().list(
                    courseId=course_id,
                    courseWorkId=cw_id,
                    userId='me'
                ).execute()
                submissions = sub_res.get('studentSubmissions', [])
                if submissions:
                    sub = submissions[0]
                    state = sub.get('state')
                    assigned_grade = sub.get('assignedGrade')

                    assignment_sub = sub.get('assignmentSubmission', {})
                    attachments = assignment_sub.get('attachments', [])

                    if state == 'RETURNED':
                        if assigned_grade is not None:
                            classroom_status = 'CALIFICADA'
                        else:
                            classroom_status = 'DEVUELTA'
                    elif state == 'TURNED_IN':
                        classroom_status = 'ENTREGADA'
                    elif attachments and state in ('NEW', 'CREATED'):
                        classroom_status = 'SUBIDA_SIN_ENTREGAR'
                    else:
                        classroom_status = 'PENDIENTE'
            except Exception:
                classroom_status = 'PENDIENTE'

            materials = cw.get('materials', [])
            attachment_links = []
            attachment_files = []

            for m in materials:
                drive_file = m.get('driveFile', {}).get('driveFile', {})
                if drive_file:
                    f_id = drive_file.get('id')
                    f_title = drive_file.get('title', 'Documento adjunto')
                    f_link = inject_authuser(drive_file.get('alternateLink', ''), user_email)

                    if f_link:
                        attachment_links.append(f"{f_title} ({f_link})")
                    if f_id:
                        attachment_files.append({'id': f_id, 'title': f_title, 'link': f_link})

                mat_link = m.get('link', {})
                if mat_link and mat_link.get('url'):
                    l_url = inject_authuser(mat_link.get('url', ''), user_email)
                    l_title = mat_link.get('title', 'Documento adjunto')
                    m_drive = re.search(r'drive\.google\.com/file/d/([a-zA-Z0-9_-]+)', l_url)
                    if m_drive:
                        d_id = m_drive.group(1)
                        if not any(af.get('id') == d_id for af in attachment_files):
                            attachment_files.append({'id': d_id, 'title': l_title, 'link': l_url})
                            attachment_links.append(f"{l_title} ({l_url})")

            full_desc = desc
            if attachment_links:
                full_desc += "\n\nArchivos adjuntos:\n" + "\n".join(attachment_links)

            cached_actions = []
            for af in attachment_files:
                af_id = af.get('id')
                if af_id:
                    c_data = get_cached_attachment_data(af_id)
                    if c_data and c_data.get('actions'):
                        cached_actions = c_data['actions']
                        break

            tasks.append({
                'id': str(cw_id),
                'course_id': course_id,
                'course_name': course_name,
                'title': title,
                'description': full_desc,
                'link': alt_link,
                'due_date': due_date_iso,
                'classroom_status': classroom_status,
                'assigned_grade': assigned_grade,
                'max_points': max_points,
                'attachment_files': attachment_files,
                'cached_actions': cached_actions
            })

    print(f"[Classroom API] Total de tareas procesadas: {len(tasks)}")
    return tasks

# Exportar con ambos nombres para compatibilidad total
fetch_tasks = get_all_tasks

def fetch_courses(creds=None, user_email=None):
    try:
        service = get_classroom_service(creds=creds)
        if not user_email and creds:
            try:
                profile = service.userProfiles().get(userId='me').execute()
                user_email = profile.get('emailAddress', '').lower().strip()
            except Exception:
                pass

        courses = get_active_courses(service)
        return [{
            'id': c.get('id'),
            'name': c.get('name', 'Materia sin nombre').replace('_', ' '),
            'section': c.get('section', ''),
            'alternateLink': inject_authuser(c.get('alternateLink', ''), user_email)
        } for c in courses]
    except Exception as e:
        print(f"Error en fetch_courses: {e}")
        return []

def get_announcements_and_alerts(creds=None, user_email=None):
    alerts = []
    try:
        service = get_classroom_service(creds=creds)
        if not user_email and creds:
            try:
                profile = service.userProfiles().get(userId='me').execute()
                user_email = profile.get('emailAddress', '').lower().strip()
            except Exception:
                pass

        courses = get_active_courses(service)

        cutoff_date = (datetime.datetime.utcnow() - datetime.timedelta(days=7)).isoformat() + "Z"

        for course in courses:
            c_id = course['id']
            c_name = course.get('name', 'Materia').replace('_', ' ')

            try:
                ann_res = service.courses().announcements().list(courseId=c_id).execute()
                announcements = ann_res.get('announcements', [])
                for a in announcements:
                    created = a.get('creationTime', '')
                    if created >= cutoff_date:
                        text = a.get('text', '').strip()
                        alerts.append({
                            'source': 'Classroom',
                            'course_name': c_name,
                            'title': f"Aviso en {c_name}",
                            'content': text,
                            'link': inject_authuser(a.get('alternateLink', ''), user_email),
                            'date': created
                        })
            except Exception:
                pass

        gmail_service = get_gmail_service(creds=creds)
        if gmail_service:
            try:
                query = "newer_than:7d (clase OR suspende OR asistencia OR aviso OR cancela OR examen OR práctica)"
                msgs_res = gmail_service.users().messages().list(userId='me', q=query, maxResults=5).execute()
                messages = msgs_res.get('messages', [])

                for m in messages:
                    msg_data = gmail_service.users().messages().get(userId='me', id=m['id'], format='snippet').execute()
                    snippet = msg_data.get('snippet', '')
                    alerts.append({
                        'source': 'Gmail',
                        'course_name': 'Correo Institucional',
                        'title': 'Aviso urgente por Correo',
                        'content': snippet,
                        'link': f"https://mail.google.com/mail/u/0/#inbox/{m['id']}",
                        'date': datetime.datetime.utcnow().isoformat()
                    })
            except Exception:
                pass

    except Exception as e:
        print(f"Error en get_announcements_and_alerts: {e}")

    return alerts
